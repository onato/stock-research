#!/usr/bin/env python3
"""{TICKER}_DCF_Model.xlsx from a built DCF: the JSON's model with live formulas.

Layout follows dcf-methods/references/owner-fcf.md section 6. Tabs in order:
Model_Info, Reconciliation (updates only), Actuals, Assumptions, DCF_Bear,
DCF_Base, DCF_Bull, Valuation. On a DCF_* tab column B is the FY0 actual and
C.. the forecast years; rows 5 (growth), 7 (EBITDA margin), 9 (SBC %) and 14
(cash tax rate) link to the Assumptions tab and every other forecast cell is
a formula, so editing a driver on Assumptions flows through to the value per
share at $B$38, the street value at $B$43 and the hurdle entry price at
$B$47. The Valuation tab probability-weights the three, and rebuilds the
required-return table and the WACC x g grid in closed form.

Rates are stored as PERCENTS (as in the JSON) and divided by 100 in the
formulas. Text cells never begin with '=' and sheet names carry no spaces:
LibreOffice's headless recalc, which the tests run, fails on either.

openpyxl writes formulas without cached values, so a freshly written book is
blank in every previewer (Finder Quick Look, GitHub, mail clients) until a
spreadsheet recalculates it. `recalc()` runs LibreOffice headless over the
file in place when it is installed, which caches every value and keeps the
formulas; `build_dcf` calls it after every write. Model_Info also carries
the methodology, data sources and scenario narratives (pitfalls.md 12): the
reader of the .xlsx never opens the JSON.
"""

from __future__ import annotations

import pathlib
import shutil
import subprocess
import tempfile
from typing import Any

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

INPUT = Font(color="0000FF")          # hardcoded driver
LINK = Font(color="008000")           # link to another tab
BOLD = Font(bold=True)
SCENARIOS = ("bear", "base", "bull")
DEFAULT_CAP = 20                      # dcf_engine.DEFAULT_CAP, for the street terminal fallback

# Assumptions tab: scalar inputs in column B, then one block per scenario.
A_PRICE, A_SHARES, A_NET_DEBT, A_BRIDGE, A_HURDLE, A_BASE_REV, A_FY0, A_CUR, A_FX = range(2, 11)
BLOCK_START = 12
BLOCK_ROWS = 14
# Offsets inside a scenario block.
O_GROWTH, O_MARGIN, O_SBC, O_LEASE, O_DA, O_CAPEX, O_TAX = range(1, 8)
O_WACC, O_TG, O_CAP, O_WC, O_WEIGHT = range(8, 13)


def _block(sc: str) -> int:
    return BLOCK_START + SCENARIOS.index(sc) * BLOCK_ROWS


def _kv(ws: Worksheet, rows: list[tuple[str, Any]], start: int = 1) -> None:
    for i, (k, v) in enumerate(rows):
        ws.cell(row=start + i, column=1, value=k)
        ws.cell(row=start + i, column=2, value=v)


def _text(v: Any) -> str | None:
    """A JSON block as one readable cell: dicts as 'key: value' lines, lists
    joined with '; '. Never starts with '=' (LibreOffice would parse it)."""
    if v is None:
        return None
    if isinstance(v, dict):
        out = "\n".join(f"{k}: {_text(x)}" for k, x in v.items())
    elif isinstance(v, list):
        out = "; ".join(str(_text(x)) for x in v)
    else:
        out = str(v)
    return ("'" + out) if out.startswith("=") else out


def _provenance_rows(dcf: dict[str, Any]) -> list[tuple[str, Any]]:
    rows: list[tuple[str, Any]] = []
    if dcf.get("valuation_philosophy") is not None:
        rows.append(("Methodology", _text(dcf["valuation_philosophy"])))
    if dcf.get("wacc_rationale") is not None:
        rows.append(("WACC rationale", _text(dcf["wacc_rationale"])))
    if dcf.get("data_sources") is not None:
        rows.append(("Data sources", _text(dcf["data_sources"])))
    raw = dcf.get("scenario_narratives")
    narratives: dict[str, Any] = raw if isinstance(raw, dict) else {}
    assumptions: dict[str, Any] = dcf.get("assumptions") or {}
    for sc in SCENARIOS:
        text = narratives.get(sc) or (assumptions.get(sc) or {}).get("narrative")
        if text:
            rows.append((f"Scenario narrative: {sc}", _text(text)))
    return rows


def recalc(path: pathlib.Path, timeout: int = 180) -> bool:
    """Recalculate the workbook in place with LibreOffice headless so every
    formula carries a cached value. Returns False (file untouched) when
    LibreOffice is not installed or the conversion fails."""
    if shutil.which("soffice") is None:
        return False
    with tempfile.TemporaryDirectory() as tmp:
        try:
            subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", tmp, str(path)],
                           check=True, capture_output=True, timeout=timeout)
        except (subprocess.CalledProcessError, subprocess.TimeoutExpired, OSError):
            return False
        out = pathlib.Path(tmp) / path.name
        if not out.exists():
            return False
        shutil.copyfile(out, path)
    return True


def _model_info(ws: Worksheet, dcf: dict[str, Any]) -> None:
    inputs = dcf.get("inputs", {})
    eng = dcf.get("engine", {})
    _kv(ws, [
        ("Ticker", dcf.get("ticker")), ("Company", dcf.get("company_name")),
        ("Valuation date", dcf.get("valuation_date")), ("Current price", dcf.get("current_price")),
        ("Price as of", inputs.get("price_as_of")), ("Balance sheet date", inputs.get("balance_sheet_date")),
        ("Model", dcf.get("model")), ("Currency (model)", inputs.get("currency")),
        ("Currency (quote)", inputs.get("quote_currency")), ("Units", inputs.get("units")),
        ("Engine", f"{eng.get('name', 'dcf_engine')} v{eng.get('version', '?')}"),
        ("Built at", eng.get("built_at")), ("Drivers file", eng.get("drivers")),
        ("Convention", ("Percents stored as percents; formulas divide by 100. Blue = input, "
                        "green = link, black = formula.")),
        *_provenance_rows(dcf),
    ])
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 100


def _reconciliation(ws: Worksheet, rec: dict[str, Any]) -> None:
    ws["A1"], ws["B1"], ws["C1"] = "Item", "Value", "Note"
    for c in ("A1", "B1", "C1"):
        ws[c].font = BOLD
    r = 2
    for label, block in (("Prior", rec.get("prior_version")), ("New", rec.get("new_version"))):
        if isinstance(block, dict):
            for k, v in block.items():
                ws.cell(row=r, column=1, value=f"{label} {k}")
                ws.cell(row=r, column=2, value=v)
                r += 1
    r += 1
    ws.cell(row=r, column=1, value="Bridge").font = BOLD
    r += 1
    for item in rec.get("bridge") or []:
        if isinstance(item, dict):
            ws.cell(row=r, column=1, value=item.get("item"))
            ws.cell(row=r, column=2, value=item.get("value"))
            ws.cell(row=r, column=3, value=item.get("note"))
            r += 1
    ws.column_dimensions["A"].width = 48


def _actuals(ws: Worksheet, history: list[dict[str, Any]]) -> None:
    if not history:
        ws["A1"] = "Period"
        ws["A2"] = "(no history supplied)"
        return
    cols = ["Period"] + [k for k in history[0] if k != "Period"]
    for j, k in enumerate(cols, 1):
        ws.cell(row=1, column=j, value=k).font = BOLD
    for i, row in enumerate(history, 2):
        for j, k in enumerate(cols, 1):
            ws.cell(row=i, column=j, value=row.get(k))


def _assumptions(ws: Worksheet, dcf: dict[str, Any], n: int) -> None:
    inputs = dcf["inputs"]
    shares = inputs["projected_shares"][-1] if inputs.get("projected_shares") else inputs["shares_outstanding"]
    cur, quote = inputs.get("currency"), inputs.get("quote_currency")
    fx = inputs.get("fx_rate") if cur and quote and cur != quote else 1.0
    bridge = sum(float(b.get("value") or 0) for b in dcf.get("equity_bridge") or [])
    ws["A1"], ws["B1"] = "Input", "Value"
    ws["A1"].font = ws["B1"].font = BOLD
    _kv(ws, [
        ("Current price (quote currency)", dcf["current_price"]),
        ("Shares, final projected (m)", shares),
        ("Net debt (m, model currency)", inputs["net_debt"]),
        ("Equity bridge total (m, + adds)", bridge),
        ("Hurdle rate (%)", float(dcf["entry_price"]["hurdle_rate"]) * 100),
        ("Base revenue FY0 (m)", inputs["base_revenue"]),
        ("FY0", inputs.get("fy0")),
        ("Currency / units", f"{cur or '?'} {inputs.get('units') or ''}".strip()),
        ("FX rate (model -> quote)", fx),
    ], start=A_PRICE)
    for r in range(A_PRICE, A_FX + 1):
        ws.cell(row=r, column=2).font = INPUT

    years = dcf["projections"]["base"]["years"]
    for sc in SCENARIOS:
        a = dcf["assumptions"][sc]
        s = _block(sc)
        ws.cell(row=s, column=1, value=sc.upper()).font = BOLD
        for j, y in enumerate(years, 3):
            ws.cell(row=s, column=j, value=y).font = BOLD
        paths = [
            (O_GROWTH, "Revenue growth (%)", a.get("growth_rates")),
            (O_MARGIN, "EBITDA margin (%)", a.get("ebitda_margin_path")),
            (O_SBC, "SBC (% of revenue)", a.get("sbc_pct_path", 0.0)),
            (O_LEASE, "Lease cost (% of revenue)", a.get("lease_cost_pct_path", 0.0)),
            (O_DA, "D&A (% of revenue)", a.get("da_pct_path", 0.0)),
            (O_CAPEX, "Capex (% of revenue)", a.get("capex_pct_path", 0.0)),
            (O_TAX, "Cash tax rate (%)", a.get("cash_tax_rate_path", 0.0)),
        ]
        for off, label, path in paths:
            ws.cell(row=s + off, column=1, value=label)
            vals = _padded(path, n)
            for j, v in enumerate(vals, 3):
                ws.cell(row=s + off, column=j, value=v).font = INPUT
        scalars = [
            (O_WACC, "WACC (%)", a["wacc"]), (O_TG, "Terminal growth (%)", a["terminal_growth"]),
            (O_CAP, "Terminal cap multiple (x FCF)", a.get("terminal_cap_multiple", DEFAULT_CAP)),
            (O_WC, "Working-capital capture (% of revenue change)", a.get("wc_capture_pct", 0.0)),
            (O_WEIGHT, "Scenario weight", dcf["probability_weighted"]["weights"][sc]),
        ]
        for off, label, v in scalars:
            ws.cell(row=s + off, column=1, value=label)
            ws.cell(row=s + off, column=2, value=v).font = INPUT
    ws.column_dimensions["A"].width = 44


def _padded(path: Any, n: int) -> list[float]:
    if isinstance(path, list) and path:
        vals = [float(x) for x in path[:n]]
        while len(vals) < n:
            vals.append(vals[-1])
        return vals
    return [float(path or 0.0)] * n


def _dcf_sheet(ws: Worksheet, dcf: dict[str, Any], sc: str, n: int) -> None:
    s = _block(sc)
    last = get_column_letter(2 + n)            # last forecast column
    ws["A1"] = f"{dcf['ticker']} owner-FCF component DCF: {sc} case (millions, {dcf['inputs'].get('currency', '')})"
    ws["A1"].font = BOLD
    ws["A3"], ws["B3"] = "Line", str(dcf["inputs"].get("fy0") or "FY0")
    for j, y in enumerate(dcf["projections"][sc]["years"], 3):
        ws.cell(row=3, column=j, value=y).font = BOLD
    ws["A4"], ws["B4"] = "Year index (discount exponent)", 0
    rows = {
        5: "Revenue growth (%)", 6: "Revenue", 7: "EBITDA margin (%)", 8: "EBITDA",
        9: "SBC (% of revenue)", 10: "SBC charge", 11: "Owner EBITDA (after SBC and leases)",
        12: "D&A", 13: "EBIT", 14: "Cash tax rate (%)", 15: "NOPAT", 16: "D&A add-back",
        17: "Capex", 18: "Working capital", 19: "Unlevered owner FCF", 20: "Discount factor",
        21: "PV of FCF", 23: "Lease cost (% of revenue)", 24: "Lease cost",
    }
    for r, label in rows.items():
        ws.cell(row=r, column=1, value=label)
    ws["B6"] = dcf["inputs"]["base_revenue"]
    ws["B6"].font = INPUT
    for i in range(n):
        c = get_column_letter(3 + i)
        p = get_column_letter(2 + i)
        f: dict[int, str] = {
            4: f"={p}4+1",
            5: f"=Assumptions!{c}{s + O_GROWTH}",
            6: f"={p}6*(1+{c}5/100)",
            7: f"=Assumptions!{c}{s + O_MARGIN}",
            8: f"={c}6*{c}7/100",
            9: f"=Assumptions!{c}{s + O_SBC}",
            10: f"=-{c}6*{c}9/100",
            11: f"={c}8+{c}10+{c}24",
            12: f"=-{c}6*Assumptions!{c}{s + O_DA}/100",
            13: f"={c}11+{c}12",
            14: f"=Assumptions!{c}{s + O_TAX}",
            15: f"={c}13-MAX(0,{c}13)*{c}14/100",
            16: f"=-{c}12",
            17: f"=-{c}6*Assumptions!{c}{s + O_CAPEX}/100",
            18: f"=({c}6-{p}6)*Assumptions!$B${s + O_WC}/100",
            19: f"={c}15+{c}16+{c}17+{c}18",
            20: f"=1/(1+$B$27/100)^{c}4",
            21: f"={c}19*{c}20",
            23: f"=Assumptions!{c}{s + O_LEASE}",
            24: f"=-{c}6*{c}23/100",
        }
        for r, formula in f.items():
            cell = ws[f"{c}{r}"]
            cell.value = formula
            if "Assumptions!" in formula:
                cell.font = LINK

    ws["A26"] = "Valuation"
    ws["A26"].font = BOLD
    val = [
        (27, "WACC (%)", f"=Assumptions!$B${s + O_WACC}"),
        (28, "Terminal growth (%)", f"=Assumptions!$B${s + O_TG}"),
        (29, "Terminal cap multiple (x FCF)", f"=Assumptions!$B${s + O_CAP}"),
        (30, "Sum of PV of FCF", f"=SUM(C21:{last}21)"),
        (31, "Terminal-year FCF", f"={last}19"),
        (32, "Gordon terminal value", "=IF(B27>B28,B31*(1+B28/100)/((B27-B28)/100),B31*B29)"),
        (33, "Cap terminal value", "=B31*B29"),
        (34, "Terminal value used: MAX(0,MIN(Gordon,cap))", "=MAX(0,MIN(B32,B33))"),
        (35, "PV of terminal value", f"=B34/(1+B27/100)^{last}4"),
        (36, "Enterprise value", "=B30+B35"),
        (37, "Equity value (EV + bridge - net debt)", f"=B36+Assumptions!$B${A_BRIDGE}-Assumptions!$B${A_NET_DEBT}"),
        (38, "Value per share (quote currency)", f"=B37/Assumptions!$B${A_SHARES}*Assumptions!$B${A_FX}"),
        (39, "Upside vs price", f"=B38/Assumptions!$B${A_PRICE}-1"),
        (41, "Street memo: PV of explicit-period SBC", f"=SUMPRODUCT(-C10:{last}10/(1+$B$27/100)^C4:{last}4)"),
        (42, "Street memo: PV of terminal SBC (Gordon)",
         f"=IF(B27>B28,-{last}10*(1+B28/100)/((B27-B28)/100),-{last}10*{DEFAULT_CAP})/(1+B27/100)^{last}4"),
        (43, "Street value per share", f"=(B37+B41+B42)/Assumptions!$B${A_SHARES}*Assumptions!$B${A_FX}"),
        (45, "Hurdle rate (%)", f"=Assumptions!$B${A_HURDLE}"),
        (46, "Terminal value at hurdle: MAX(0,MIN(Gordon@hurdle,cap))",
         "=MAX(0,MIN(IF(B45>B28,B31*(1+B28/100)/((B45-B28)/100),B31*B29),B31*B29))"),
        (47, "Hurdle entry price (quote currency)",
         (f"=(SUMPRODUCT(C19:{last}19/(1+B45/100)^C4:{last}4)+B46/(1+B45/100)^{last}4"
          f"+Assumptions!$B${A_BRIDGE}-Assumptions!$B${A_NET_DEBT})/Assumptions!$B${A_SHARES}*Assumptions!$B${A_FX}")),
    ]
    for r, label, formula in val:
        ws.cell(row=r, column=1, value=label)
        cell = ws.cell(row=r, column=2, value=formula)
        if "Assumptions!" in formula and formula.count("!") == 1 and "SUMPRODUCT" not in formula:
            cell.font = LINK
    ws.column_dimensions["A"].width = 52


def _closed_form(rate_ref: str, g_ref: str, last: str) -> str:
    """Base-case value per share at an arbitrary rate and terminal growth,
    rebuilding the terminal value at that rate (owner-fcf.md section 6)."""
    fcf = f"DCF_Base!$C$19:${last}$19"
    idx = f"DCF_Base!$C$4:${last}$4"
    fcf_n = f"DCF_Base!${last}$19"
    cap = "DCF_Base!$B$29"
    gordon = f"IF({rate_ref}>{g_ref},{fcf_n}*(1+{g_ref}/100)/(({rate_ref}-{g_ref})/100),{fcf_n}*{cap})"
    return (f"=(SUMPRODUCT({fcf}/(1+{rate_ref}/100)^{idx})+MAX(0,MIN({gordon},{fcf_n}*{cap}))"
            f"/(1+{rate_ref}/100)^DCF_Base!${last}$4+Assumptions!$B${A_BRIDGE}-Assumptions!$B${A_NET_DEBT})"
            f"/Assumptions!$B${A_SHARES}*Assumptions!$B${A_FX}")


def _valuation(ws: Worksheet, dcf: dict[str, Any], n: int) -> None:
    last = get_column_letter(2 + n)
    ws["A1"] = f"{dcf['ticker']} valuation summary"
    ws["A1"].font = BOLD
    heads = ["Scenario", "Value per share", "Upside vs price", "Weight", "Street value per share",
             "Hurdle entry price"]
    for j, h in enumerate(heads, 1):
        ws.cell(row=3, column=j, value=h).font = BOLD
    for i, sc in enumerate(SCENARIOS):
        r = 4 + i
        tab = f"DCF_{sc.capitalize()}"
        ws.cell(row=r, column=1, value=sc.capitalize())
        ws.cell(row=r, column=2, value=f"={tab}!$B$38").font = LINK
        ws.cell(row=r, column=3, value=f"=B{r}/Assumptions!$B${A_PRICE}-1")
        ws.cell(row=r, column=4, value=f"=Assumptions!$B${_block(sc) + O_WEIGHT}").font = LINK
        ws.cell(row=r, column=5, value=f"={tab}!$B$43").font = LINK
        ws.cell(row=r, column=6, value=f"={tab}!$B$47").font = LINK
    _kv(ws, [
        ("Probability-weighted value per share", "=SUMPRODUCT(B4:B6,D4:D6)"),
        ("Street weighted value per share", "=SUMPRODUCT(E4:E6,D4:D6)"),
        ("Margin of safety (1 - price/weighted)", f"=1-Assumptions!$B${A_PRICE}/B8"),
        ("Band position (%; 0 = house, 100 = street)",
         f'=IF(ABS(B9-B8)>0.01*ABS(B8),(Assumptions!$B${A_PRICE}-B8)/(B9-B8)*100,"n/a")'),
        ("Weighted 15% hurdle entry price", "=SUMPRODUCT(F4:F6,D4:D6)"),
    ], start=8)

    ws["A14"] = "Required-return table (base case; each row rebuilds its own terminal value)"
    ws["A14"].font = BOLD
    ws["C14"] = "Rate (%)"
    rrt = dcf["required_return_table"]
    wacc = dcf["assumptions"]["base"]["wacc"]
    r = 15
    for rate in rrt["returns"]:
        label = f"Value per share at {rate:g}%" + (" (WACC)" if float(rate) == float(wacc) else "")
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=3, value=float(rate)).font = INPUT
        ws.cell(row=r, column=2, value=_closed_form(f"$C${r}", "DCF_Base!$B$28", last))
        r += 1

    r += 1
    ws.cell(row=r, column=1, value="Sensitivity: value per share, rows WACC (%) x columns terminal growth (%)").font = BOLD
    r += 1
    sens = dcf["sensitivity"]
    for j, g in enumerate(sens["terminal_growth_range"], 2):
        ws.cell(row=r, column=j, value=float(g)).font = INPUT
    head_row = r
    for i, w in enumerate(sens["wacc_range"], 1):
        ws.cell(row=r + i, column=1, value=float(w)).font = INPUT
        for j in range(len(sens["terminal_growth_range"])):
            col = get_column_letter(2 + j)
            ws.cell(row=r + i, column=2 + j,
                    value=_closed_form(f"$A${r + i}", f"${col}${head_row}", last))
    ws.column_dimensions["A"].width = 60


def write(dcf: dict[str, Any], path: pathlib.Path, history: list[dict[str, Any]] | None = None) -> pathlib.Path:
    """Write the workbook for a built DCF. `history` (rows of the Metrics
    CSV, dicts keyed by CSV header) fills the Actuals tab when given."""
    n = len(dcf["projections"]["base"]["fcf"])
    wb = openpyxl.Workbook()
    info = wb.active
    assert isinstance(info, Worksheet)
    _model_info(info, dcf)
    info.title = "Model_Info"
    rec = dcf.get("reconciliation")
    if isinstance(rec, dict):
        _reconciliation(wb.create_sheet("Reconciliation"), rec)
    _actuals(wb.create_sheet("Actuals"), history or [])
    _assumptions(wb.create_sheet("Assumptions"), dcf, n)
    for sc in SCENARIOS:
        _dcf_sheet(wb.create_sheet(f"DCF_{sc.capitalize()}"), dcf, sc, n)
    _valuation(wb.create_sheet("Valuation"), dcf, n)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
