"""scripts/build_dcf.py -- the CLI over dcf_engine + dcf_workbook.

`make build-dcf TICKER=X` reads Reports/X_Drivers.json and writes X_DCF.json
and X_DCF_Model.xlsx; `--check` re-derives an existing X_DCF.json from its
own assumptions. The generated JSON must pass the dashboard's node verify
gate, which is the only executable definition of "engine-compatible".
"""

import json
import shutil
from pathlib import Path

import build_dashboard as bd
import build_dcf
import pytest

FIX = Path(__file__).parent / "fixtures" / "dcf"
DASH_FIX = Path(__file__).parent / "fixtures" / "dashboard"


@pytest.fixture
def repo(tmp_path, monkeypatch):
    reports = tmp_path / "research" / "APA.AX" / "Reports"
    reports.mkdir(parents=True)
    shutil.copy(FIX / "APA_Drivers.json", reports / "APA.AX_Drivers.json")
    (reports / "APA.AX_Metrics.csv").write_text(
        "Period,Revenue,EBITDA,FreeCashFlow,Units,Currency\n"
        "FY2025,3179,1894,1083,millions,AUD\n"
        "H1 FY2026,1614,1046,556,millions,AUD\n"
        "FY2026,2977,2095,1118,millions,AUD\n")
    monkeypatch.setattr(build_dcf, "REPO", tmp_path)
    return tmp_path


class TestBuild:
    def test_writes_json_and_workbook(self, repo, capsys):
        assert build_dcf.main(["APA.AX"]) == 0
        reports = repo / "research" / "APA.AX" / "Reports"
        dcf = json.loads((reports / "APA.AX_DCF.json").read_text())
        assert dcf["probability_weighted"]["weighted_iv"] == 7.33
        assert (reports / "APA.AX_DCF_Model.xlsx").stat().st_size > 10000
        out = capsys.readouterr().out
        assert "weighted_iv 7.33" in out
        assert "base 8.0" in out
        assert "engine check: base:ok bull:ok bear:ok" in out

    def test_actuals_tab_holds_the_annual_rows_only(self, repo):
        import openpyxl
        build_dcf.main(["APA.AX"])
        ws = openpyxl.load_workbook(repo / "research" / "APA.AX" / "Reports" / "APA.AX_DCF_Model.xlsx")["Actuals"]
        assert [ws.cell(row=r, column=1).value for r in (2, 3)] == ["FY2025", "FY2026"]
        assert ws.cell(row=4, column=1).value is None

    def test_bad_drivers_name_the_fields_and_write_nothing(self, repo, capsys):
        reports = repo / "research" / "APA.AX" / "Reports"
        d = json.loads((reports / "APA.AX_Drivers.json").read_text())
        del d["inputs"]["net_debt"]
        d["assumptions"]["bear"]["capex_pct"] = 20
        (reports / "APA.AX_Drivers.json").write_text(json.dumps(d))
        assert build_dcf.main(["APA.AX"]) == 2
        err = capsys.readouterr().err
        assert "inputs.net_debt" in err
        assert "assumptions.bear.capex_pct" in err
        assert not (reports / "APA.AX_DCF.json").exists()

    def test_missing_drivers_file(self, repo, capsys):
        assert build_dcf.main(["XYZ.NZ"]) == 2
        assert "XYZ.NZ_Drivers.json" in capsys.readouterr().err

    def test_a_prior_dcf_is_kept_as_a_backup(self, repo):
        reports = repo / "research" / "APA.AX" / "Reports"
        (reports / "APA.AX_DCF.json").write_text('{"old": true}')
        build_dcf.main(["APA.AX"])
        assert (reports / "APA.AX_DCF.prev.json").read_text() == '{"old": true}'


class TestCheck:
    def test_check_reports_each_scenario(self, repo, capsys):
        build_dcf.main(["APA.AX"])
        capsys.readouterr()
        assert build_dcf.main(["APA.AX", "--check"]) == 0
        assert "APA.AX  base:ok bull:ok bear:ok" in capsys.readouterr().out

    def test_check_fails_on_a_mismatch(self, repo, capsys):
        build_dcf.main(["APA.AX"])
        p = repo / "research" / "APA.AX" / "Reports" / "APA.AX_DCF.json"
        d = json.loads(p.read_text())
        d["valuation"]["bull"]["intrinsic_value"] = 12.0
        p.write_text(json.dumps(d))
        assert build_dcf.main(["APA.AX", "--check"]) == 1
        assert "bull:mismatch" in capsys.readouterr().out


@pytest.mark.skipif(shutil.which("node") is None, reason="node not installed")
class TestDashboardGate:
    def test_generated_json_passes_the_slider_engine_gate(self, repo):
        build_dcf.main(["APA.AX"])
        dcf = json.loads((repo / "research" / "APA.AX" / "Reports" / "APA.AX_DCF.json").read_text())
        spec = json.loads((DASH_FIX / "TEST_DashboardSpec.json").read_text())
        analysis = json.loads((DASH_FIX / "TEST_Analysis.json").read_text())
        page = bd.render("APA.AX", spec, (DASH_FIX / "TEST_Metrics.csv").read_text(), analysis, dcf)
        status = bd.engine_status(page)
        assert status is not None
        assert {sc: status[sc]["ok"] for sc in ("base", "bull", "bear")} == {"base": True, "bull": True, "bear": True}
        assert status["base"]["family"] == "component"
