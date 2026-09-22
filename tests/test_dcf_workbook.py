"""scripts/dcf_workbook.py -- the DCF_Model.xlsx from a built DCF.

The dcf-analyst used to write a bespoke ~22k-char openpyxl generator per
ticker. The workbook is the same model as the JSON with live formulas
(owner-fcf.md section 6), so it is generated from the engine's output. The
recalc test is the reference's own verification step: LibreOffice headless,
zero '#' cells, scenario values and the weighted figure agree with the JSON.
"""

import json
import shutil
import subprocess
from pathlib import Path

import dcf_engine
import dcf_workbook
import openpyxl
import pytest

FIX = Path(__file__).parent / "fixtures" / "dcf"


@pytest.fixture(scope="module")
def dcf():
    return dcf_engine.build(json.loads((FIX / "APA_Drivers.json").read_text()))


@pytest.fixture(scope="module")
def book_path(dcf, tmp_path_factory):
    out = tmp_path_factory.mktemp("wb") / "APA.AX_DCF_Model.xlsx"
    dcf_workbook.write(dcf, out, history=[
        {"Period": "FY2025", "Revenue": 3179.0, "EBITDA": 1894.0, "FreeCashFlow": 1083.0},
        {"Period": "FY2026", "Revenue": 2977.0, "EBITDA": 2095.0, "FreeCashFlow": 1118.0},
    ])
    return out


@pytest.fixture(scope="module")
def book(book_path):
    return openpyxl.load_workbook(book_path)


class TestStructure:
    def test_tabs_in_the_references_order(self, book):
        assert book.sheetnames == ["Model_Info", "Actuals", "Assumptions", "DCF_Bear", "DCF_Base",
                                   "DCF_Bull", "Valuation"]

    def test_reconciliation_tab_only_on_updates(self, dcf, tmp_path):
        d = dict(dcf)
        d["reconciliation"] = {"prior_version": {"valuation_date": "2026-09-21", "weighted": 7.33},
                               "new_version": {"weighted": 7.33},
                               "bridge": [{"item": "Prior weighted IV", "value": 7.33},
                                          {"item": "Price refresh", "value": 0.0}]}
        out = tmp_path / "x.xlsx"
        dcf_workbook.write(d, out)
        assert openpyxl.load_workbook(out).sheetnames[1] == "Reconciliation"

    def test_sheet_names_have_no_spaces(self, book):
        assert all(" " not in n for n in book.sheetnames)

    def test_text_cells_never_begin_with_equals(self, book):
        # LibreOffice parses such a label as a formula and the recalc errors out.
        for ws in book.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith("=") and c.data_type != "f":
                        pytest.fail(f"{ws.title}!{c.coordinate} text begins with '='")

    def test_the_row_map(self, book):
        ws = book["DCF_Base"]
        labels = {r: ws.cell(row=r, column=1).value for r in (5, 6, 7, 8, 9, 10, 12, 13, 14, 15, 17, 19, 20, 21)}
        assert labels[5].lower().startswith("revenue growth")
        assert labels[6] == "Revenue"
        assert labels[19].lower().startswith("unlevered owner fcf")
        assert labels[21].lower().startswith("pv")
        assert ws["B6"].value == 2977.0                       # FY0 actual, a value
        assert str(ws["C6"].value).startswith("=")            # forecasts are formulas
        assert str(ws["C19"].value).startswith("=")
        assert str(ws["B38"].value).startswith("=")           # value per share
        assert str(ws["B43"].value).startswith("=")           # street value per share

    def test_only_driver_rows_hold_hardcoded_numbers(self, book):
        ws = book["DCF_Base"]
        for r in range(6, 22):
            if r in (5, 7, 9, 14):
                continue
            for col in "CDEFGHIJKL":
                v = ws[f"{col}{r}"].value
                assert v is None or str(v).startswith("="), f"{col}{r} = {v!r}"

    def test_drivers_link_to_the_assumptions_tab(self, book):
        assert "Assumptions!" in str(book["DCF_Base"]["C5"].value)
        assert "Assumptions!" in str(book["DCF_Bull"]["C7"].value)

    def test_valuation_tab_weights_and_band(self, book):
        ws = book["Valuation"]
        cells = [str(c.value) for row in ws.iter_rows() for c in row if c.value is not None]
        assert any("SUMPRODUCT" in c for c in cells)
        assert any("DCF_Base!$B$38" in c for c in cells)
        assert any("DCF_Bear!$B$43" in c for c in cells)

    def test_actuals_come_from_history(self, book):
        ws = book["Actuals"]
        assert ws["A1"].value == "Period"
        assert ws["A3"].value == "FY2026"
        assert ws["B3"].value == 2977.0


@pytest.fixture(scope="module")
def values(book_path, tmp_path_factory):
    """The workbook after a LibreOffice headless recalc, values only."""
    out = tmp_path_factory.mktemp("recalc")
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(out),
                    str(book_path)], check=True, capture_output=True, timeout=120)
    return openpyxl.load_workbook(out / book_path.name, data_only=True)


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
class TestRecalc:
    def test_zero_formula_errors(self, values):
        bad = [f"{ws.title}!{c.coordinate}={c.value}" for ws in values.worksheets
               for row in ws.iter_rows() for c in row
               if isinstance(c.value, str) and c.value.startswith("#")]
        assert bad == []

    def test_scenario_values_agree_with_the_json(self, values, dcf):
        for sc in ("Bear", "Base", "Bull"):
            ws = values[f"DCF_{sc}"]
            assert ws["B38"].value == pytest.approx(dcf["valuation"][sc.lower()]["intrinsic_value"], abs=0.01)
            assert ws["B43"].value == pytest.approx(dcf["valuation"][sc.lower()]["street_intrinsic_value"], abs=0.01)
            assert ws["C19"].value == pytest.approx(dcf["projections"][sc.lower()]["fcf"][0], abs=0.1)

    def test_weighted_entry_and_required_return_agree(self, values, dcf):
        ws = values["Valuation"]
        got = {str(ws.cell(row=r, column=1).value): ws.cell(row=r, column=2).value
               for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=1).value}
        assert got["Probability-weighted value per share"] == pytest.approx(dcf["probability_weighted"]["weighted_iv"], abs=0.01)
        assert got["Weighted 15% hurdle entry price"] == pytest.approx(dcf["entry_price"]["weighted_entry_price"], abs=0.01)
        assert got["Value per share at 15%"] == pytest.approx(dcf["entry_price"]["base"]["entry_price"], abs=0.01)
        assert got["Value per share at 10%"] == pytest.approx(dcf["required_return_table"]["value_per_share"][2], abs=0.01)
