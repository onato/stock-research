"""scripts/repair_xlsx.py -- text labels that openpyxl stored as formulas.

A hand-written workbook whose label began with '= ' (e.g. '= Owner FCF')
was saved by openpyxl as <f> Owner FCF</f>. Excel repairs it on open;
Numbers refuses the file outright. The repair turns such cells back into
strings and leaves real formulas, even ones with a leading space, alone.
"""

import subprocess
import sys
from pathlib import Path

import openpyxl
import repair_xlsx


def _book(path: Path, **cells: object) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    for ref, v in cells.items():
        ws[ref] = v
    wb.save(path)
    return path


def test_text_stored_as_formula_becomes_a_string(tmp_path):
    p = _book(tmp_path / "x.xlsx", A1="= Owner FCF (NZ$M)", A2=5, B1="=A2*2",
              C1="= 11.5% base WACC (10.5% bull / 13.0% bear)")
    assert repair_xlsx.repair(p) == 2
    ws = openpyxl.load_workbook(p).active
    assert ws["A1"].data_type == "s"
    assert ws["A1"].value == "= Owner FCF (NZ$M)"
    assert ws["C1"].value == "= 11.5% base WACC (10.5% bull / 13.0% bear)"
    assert ws["B1"].data_type == "f"
    assert ws["B1"].value == "=A2*2"
    assert ws["A2"].value == 5


def test_real_formula_with_a_leading_space_is_left_alone(tmp_path):
    p = _book(tmp_path / "x.xlsx", A1="= A2*2", A2=5, B1="= SUM('DCF Model'!A1:A2)")
    assert repair_xlsx.repair(p) == 0
    ws = openpyxl.load_workbook(p).active
    assert ws["A1"].data_type == "f"
    assert ws["B1"].data_type == "f"


def test_clean_file_is_not_rewritten(tmp_path):
    p = _book(tmp_path / "x.xlsx", A1="Owner FCF", B1="=A2*2", A2=5)
    before = p.read_bytes()
    assert repair_xlsx.repair(p) == 0
    assert p.read_bytes() == before


def test_cell_style_survives(tmp_path):
    p = tmp_path / "x.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "= Owner FCF"
    ws["A1"].font = openpyxl.styles.Font(bold=True)
    wb.save(p)
    repair_xlsx.repair(p)
    assert openpyxl.load_workbook(p).active["A1"].font.bold is True


def test_cli_check_reports_without_writing(tmp_path):
    p = _book(tmp_path / "x.xlsx", A1="= Owner FCF")
    before = p.read_bytes()
    r = subprocess.run([sys.executable, "scripts/repair_xlsx.py", "--check", str(p)],
                       capture_output=True, text=True, check=False)
    assert r.returncode == 1
    assert "1 text-as-formula cell" in r.stdout
    assert p.read_bytes() == before
    r = subprocess.run([sys.executable, "scripts/repair_xlsx.py", str(p)], capture_output=True, text=True, check=False)
    assert r.returncode == 0
    assert repair_xlsx.repair(p) == 0


# openpyxl without lxml (the system python3) serialises docProps/core.xml with
# the namespaces declared on the child elements and no dc: element at all.
# Numbers refuses such a file (WISE.L_DCF_Model.xlsx, 2026-09-22); with lxml the
# same library declares every namespace on the root and adds dc:creator.
ET_CORE = (
    '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties">'
    '<dcterms:created xmlns:dcterms="http://purl.org/dc/terms/" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance"'
    ' xsi:type="dcterms:W3CDTF">2026-08-01T06:13:41Z</dcterms:created>'
    '<dcterms:modified xmlns:dcterms="http://purl.org/dc/terms/" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance"'
    ' xsi:type="dcterms:W3CDTF">2026-08-27T22:25:21Z</dcterms:modified></cp:coreProperties>'
)


def _with_core(path: Path, core: str) -> Path:
    import zipfile
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as dst:
        for n in src.namelist():
            dst.writestr(n, core.encode() if n == "docProps/core.xml" else src.read(n))
    tmp.replace(path)
    return path


def test_core_properties_get_root_namespaces_and_a_creator(tmp_path):
    import zipfile
    p = _with_core(_book(tmp_path / "x.xlsx", A1="Owner FCF"), ET_CORE)
    assert repair_xlsx.repair(p) == 1
    core = zipfile.ZipFile(p).read("docProps/core.xml").decode()
    root = core[:core.index(">") + 1]
    for ns in ("xmlns:cp=", 'xmlns:dc="http://purl.org/dc/elements/1.1/"', "xmlns:dcterms=", "xmlns:xsi="):
        assert ns in root
    assert "<dc:creator>" in core
    assert "2026-08-01T06:13:41Z" in core
    assert "2026-08-27T22:25:21Z" in core
    assert "xmlns:dcterms" not in core[len(root):]
    assert openpyxl.load_workbook(p).active["A1"].value == "Owner FCF"
    assert repair_xlsx.repair(p) == 0


def test_lxml_style_core_is_left_alone(tmp_path):
    import zipfile
    p = _book(tmp_path / "x.xlsx", A1="Owner FCF")
    core = zipfile.ZipFile(p).read("docProps/core.xml").decode()
    assert 'xmlns:dc="http://purl.org/dc/elements/1.1/"' in core[:core.index(">")]
    assert repair_xlsx.repair(p) == 0


def test_cli_reports_the_core_fix(tmp_path):
    p = _with_core(_book(tmp_path / "x.xlsx", A1="Owner FCF"), ET_CORE)
    r = subprocess.run([sys.executable, "scripts/repair_xlsx.py", "--check", str(p)],
                       capture_output=True, text=True, check=False)
    assert r.returncode == 1
    assert "core.xml" in r.stdout


# xlsxwriter-style: XML declaration first, every namespace on the root, empty
# creator. Numbers opens these; the repair must not touch them.
DECL_CORE = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    '<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties"'
    ' xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/"'
    ' xmlns:dcmitype="http://purl.org/dc/dcmitype/" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
    '<dc:creator></dc:creator><cp:lastModifiedBy></cp:lastModifiedBy>'
    '<dcterms:created xsi:type="dcterms:W3CDTF">2026-06-27T06:13:25Z</dcterms:created>'
    '<dcterms:modified xsi:type="dcterms:W3CDTF">2026-06-27T06:13:25Z</dcterms:modified></cp:coreProperties>'
)


def test_core_with_xml_declaration_and_root_namespaces_is_left_alone(tmp_path):
    p = _with_core(_book(tmp_path / "x.xlsx", A1="Owner FCF"), DECL_CORE)
    before = p.read_bytes()
    assert repair_xlsx.repair(p) == 0
    assert p.read_bytes() == before
